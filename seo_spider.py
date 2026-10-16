#!/usr/bin/env python3
"""
SilentRat engine - full-featured, advanced SEO crawler (no URL limit).

Engine used by the GUI (silentrat.py) and as a CLI.

Core features:
  - Unlimited multithreaded HTML crawling, crawl depth, inlinks
  - Status codes, redirects, response time, page size
  - Title, meta description, meta keywords, H1, H2
  - Canonical + diagnostics, directives (meta robots, X-Robots), indexability
  - hreflang, rel next/prev, lang, viewport, charset, word count, text ratio
  - Internal/external links and outlinks, CSS/JS/image resources with status
  - Mixed content, duplicates, robots.txt, include/exclude regex

Advanced features:
  - JavaScript rendering (headless Chrome/Edge via Selenium)
  - Custom extraction (CSS / XPath / Regex, up to 5 extractors)
  - Structured data (JSON-LD / Schema.org, microdata)
  - Exact duplicate content (hash)
  - Core Web Vitals measured locally (optional PageSpeed API in the CLI)
  - Crawl modes: spider / URL list / XML sitemap
  - CSV/XLSX export and sitemap.xml generation
"""

import argparse
import csv
import gzip
import hashlib
import json
import re
import threading
import time
import urllib.robotparser
from collections import deque
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, asdict
from urllib.parse import urljoin, urlparse, urldefrag
from xml.sax.saxutils import escape as xml_escape

import requests
from bs4 import BeautifulSoup

DEFAULT_UA = (
    "Mozilla/5.0 (compatible; SilentRat/3.0; +https://example.com/bot) "
    "Python-requests"
)
MAX_CUSTOM = 5


@dataclass
class PageResult:
    url: str = ""
    status_code: int = 0
    content_type: str = ""
    indexable: str = ""
    indexability_status: str = ""
    title: str = ""
    title_len: int = 0
    meta_description: str = ""
    meta_desc_len: int = 0
    meta_keywords: str = ""
    meta_robots: str = ""
    x_robots: str = ""
    h1: str = ""
    h1_count: int = 0
    h2: str = ""
    h2_count: int = 0
    canonical: str = ""
    canonical_status: str = ""
    hreflang: str = ""
    hreflang_count: int = 0
    rel_next: str = ""
    rel_prev: str = ""
    lang: str = ""
    viewport: str = ""
    charset: str = ""
    word_count: int = 0
    text_ratio: float = 0.0
    size_kb: float = 0.0
    images: int = 0
    images_no_alt: int = 0
    internal_links: int = 0
    external_links: int = 0
    outlinks: int = 0
    inlinks: int = 0
    is_https: str = ""
    mixed_content: int = 0
    structured_data: str = ""
    structured_data_count: int = 0
    rendered: str = "No"
    content_hash: str = ""
    custom_1: str = ""
    custom_2: str = ""
    custom_3: str = ""
    custom_4: str = ""
    custom_5: str = ""
    ps_performance: str = ""
    ps_fcp: str = ""
    ps_lcp: str = ""
    ps_cls: str = ""
    ps_ttfb: str = ""
    ps_load: str = ""
    redirect_to: str = ""
    redirect_chain: str = ""
    response_ms: int = 0
    crawl_depth: int = 0
    error: str = ""


@dataclass
class ResourceResult:
    url: str = ""
    type: str = ""
    status_code: int = 0
    source: str = ""


class JSRenderer:
    """Headless renderer (Chrome, falling back to Edge). One driver per thread."""
    def __init__(self, wait=2.0, user_agent=DEFAULT_UA):
        self.wait = wait
        self.user_agent = user_agent
        self._local = threading.local()
        self._drivers = []
        self._lock = threading.Lock()
        self.available = True
        self.backend = None

    def _make_driver(self):
        from selenium import webdriver
        last_err = None
        # Chrome
        try:
            opts = webdriver.ChromeOptions()
            opts.add_argument("--headless=new")
            opts.add_argument("--disable-gpu")
            opts.add_argument("--no-sandbox")
            opts.add_argument("--window-size=1366,768")
            opts.add_argument("--log-level=3")
            opts.add_argument(f"--user-agent={self.user_agent}")
            opts.add_experimental_option("excludeSwitches", ["enable-logging"])
            d = webdriver.Chrome(options=opts)
            self.backend = "Chrome"
            return d
        except Exception as e:
            last_err = e
        # Edge
        try:
            opts = webdriver.EdgeOptions()
            opts.add_argument("--headless=new")
            opts.add_argument("--disable-gpu")
            opts.add_argument("--no-sandbox")
            opts.add_argument("--window-size=1366,768")
            opts.add_argument("--log-level=3")
            opts.add_argument(f"--user-agent={self.user_agent}")
            d = webdriver.Edge(options=opts)
            self.backend = "Edge"
            return d
        except Exception as e:
            last_err = e
        raise RuntimeError(f"Could not start a headless browser: {last_err}")

    def get_driver(self):
        d = getattr(self._local, "driver", None)
        if d is None:
            d = self._make_driver()
            self._local.driver = d
            with self._lock:
                self._drivers.append(d)
        return d

    _VITALS_JS = r"""
        try {
          const nav = performance.getEntriesByType('navigation')[0] || {};
          let fcp = '';
          const f = performance.getEntriesByName('first-contentful-paint')[0];
          if (f) fcp = (f.startTime/1000).toFixed(2) + ' s';
          let lcp = '';
          const L = performance.getEntriesByType('largest-contentful-paint');
          if (L.length) lcp = (L[L.length-1].startTime/1000).toFixed(2) + ' s';
          let cls = 0;
          for (const e of performance.getEntriesByType('layout-shift'))
            if (!e.hadRecentInput) cls += e.value;
          let ttfb = '';
          if (nav.responseStart) ttfb = Math.round(nav.responseStart) + ' ms';
          let load = '';
          if (nav.loadEventEnd) load = (nav.loadEventEnd/1000).toFixed(2) + ' s';
          return {fcp: fcp, lcp: lcp, cls: cls.toFixed(3), ttfb: ttfb, load: load};
        } catch (e) { return {}; }
    """

    def render(self, url, timeout=25, measure=False):
        d = self.get_driver()
        d.set_page_load_timeout(timeout)
        d.get(url)
        if self.wait:
            time.sleep(self.wait)
        vitals = {}
        if measure:
            try:
                vitals = d.execute_script(self._VITALS_JS) or {}
            except Exception:
                vitals = {}
        return d.page_source, vitals

    def quit_all(self):
        with self._lock:
            for d in self._drivers:
                try:
                    d.quit()
                except Exception:
                    pass
            self._drivers = []


GZIP_MAGIC = bytes([0x1F, 0x8B])


class SeoSpider:
    def __init__(self, start_url, max_urls=1_000_000, delay=0.0, threads=5,
                 user_agent=DEFAULT_UA, respect_robots=True, timeout=20,
                 include_external_status=True, check_resources=True,
                 include_pattern="", exclude_pattern="",
                 mode="spider", seed_urls=None,
                 render_js=False, render_wait=2.0,
                 extractors=None):
        parsed = urlparse(start_url)
        self.start_url = start_url
        self.scheme = parsed.scheme or "https"
        self.domain = parsed.netloc.lower()
        self.max_urls = max_urls
        self.delay = delay
        self.timeout = timeout
        self.include_external_status = include_external_status
        self.check_resources = check_resources
        self.include_re = re.compile(include_pattern) if include_pattern else None
        self.exclude_re = re.compile(exclude_pattern) if exclude_pattern else None
        self.mode = mode
        self.seed_urls = seed_urls or []
        self.follow_links = (mode == "spider")
        self.extractors = (extractors or [])[:MAX_CUSTOM]

        # JS rendering: one driver per thread is expensive, so cap the thread count.
        self.render_js = render_js
        self.renderer = JSRenderer(render_wait, user_agent) if render_js else None
        self.threads = min(threads, 3) if render_js else threads

        self.session = requests.Session()
        self.session.headers.update({"User-Agent": user_agent})

        self.results = {}
        self.resources = {}
        self.inlink_counter = {}
        self.seen = set()
        self.queue = deque()
        self.external_checked = {}
        self._resource_urls = {}
        self.lock = threading.Lock()

        self.on_result = None
        self.on_progress = None
        self.on_status = None
        self.stop_event = threading.Event()

        self.robots = None
        if respect_robots:
            self._load_robots(user_agent)

    # ---------- robots ----------
    def _load_robots(self, ua):
        robots_url = f"{self.scheme}://{self.domain}/robots.txt"
        rp = urllib.robotparser.RobotFileParser()
        try:
            resp = self.session.get(robots_url, timeout=self.timeout)
            if resp.status_code == 200:
                rp.parse(resp.text.splitlines())
                self.robots = rp
        except Exception:
            pass

    def _allowed(self, url, ua):
        if self.robots is None:
            return True
        try:
            return self.robots.can_fetch(ua, url)
        except Exception:
            return True

    # ---------- helpers ----------
    def _same_domain(self, url):
        net = urlparse(url).netloc.lower()
        base = self.domain[4:] if self.domain.startswith("www.") else self.domain
        net_base = net[4:] if net.startswith("www.") else net
        return net_base == base

    def _normalize(self, url):
        url, _ = urldefrag(url)
        if url.endswith("/") and len(urlparse(url).path) > 1:
            url = url.rstrip("/")
        return url

    def _passes_filters(self, url):
        if self.exclude_re and self.exclude_re.search(url):
            return False
        if self.include_re and not self.include_re.search(url):
            return False
        return True

    # ---------- input sitemap ----------
    SITEMAP_CANDIDATES = ("/sitemap_index.xml", "/sitemap.xml", "/wp-sitemap.xml",
                          "/sitemap.xml.gz", "/sitemap-index.xml")

    def _status(self, msg):
        if self.on_status:
            self.on_status(msg)
        else:
            print(f"[sitemap] {msg}")

    def _fetch_sitemap_doc(self, url):
        """Download a sitemap (supports .gz). Return the XML soup, or None if it is not a sitemap."""
        try:
            resp = self.session.get(url, timeout=self.timeout, allow_redirects=True)
        except requests.exceptions.RequestException as e:
            self._status(f"Sitemap unreachable {url}: {str(e)[:80]}")
            return None
        if resp.status_code != 200:
            self._status(f"Sitemap {url} returned HTTP {resp.status_code}")
            return None
        data = resp.content
        if data[:2] == GZIP_MAGIC:          # gzip without Content-Encoding (.xml.gz)
            try:
                data = gzip.decompress(data)
            except OSError:
                return None
        head = data[:2048].lower()
        if b"<urlset" not in head and b"<sitemapindex" not in head:
            return None                        # HTML or something else: not a sitemap
        return BeautifulSoup(data, "xml")

    def discover_sitemaps(self, base_url):
        """Find the site's sitemaps: Sitemap: directives in robots.txt, then common paths."""
        parsed = urlparse(base_url)
        root = f"{parsed.scheme or 'https'}://{parsed.netloc}"
        found = []
        try:
            resp = self.session.get(root + "/robots.txt", timeout=self.timeout)
            if resp.status_code == 200:
                for line in resp.text.splitlines():
                    if line.strip().lower().startswith("sitemap:"):
                        sm = line.split(":", 1)[1].strip()
                        if sm and sm not in found:
                            found.append(sm)
        except requests.exceptions.RequestException:
            pass
        if not found:
            found = [root + c for c in self.SITEMAP_CANDIDATES]
        return found

    def load_sitemap_urls(self, sitemap_url, _depth=0, _visited=None):
        """Download a sitemap.xml (or index, recursively, .gz) and return the page <loc> URLs.
        If the URL is not a sitemap (e.g. the homepage), auto-discover via robots.txt / common paths."""
        top = _visited is None
        visited = set() if top else _visited
        urls = []
        if _depth > 5 or sitemap_url in visited:
            return urls
        visited.add(sitemap_url)
        soup = self._fetch_sitemap_doc(sitemap_url)
        if soup is None:
            if top:
                self._status(f"{sitemap_url} is not an XML sitemap; looking for the site's sitemaps...")
                for cand in self.discover_sitemaps(sitemap_url):
                    got = self.load_sitemap_urls(cand, 1, visited)
                    urls.extend(got)
                    # from robots.txt load them all; for common paths the first valid one is enough
                    if got and cand.endswith(self.SITEMAP_CANDIDATES):
                        break
                urls = list(dict.fromkeys(urls))
                self._status(f"Sitemap: {len(urls)} URLs found")
            return urls
        index = soup.find("sitemapindex")
        if index is not None:
            for sm in index.find_all("sitemap", recursive=False):
                loc = sm.find("loc", recursive=False)
                if loc and loc.text.strip():
                    child = loc.text.strip()
                    self._status(f"Reading {child} ...")
                    urls.extend(self.load_sitemap_urls(child, _depth + 1, visited))
        else:
            urlset = soup.find("urlset")
            for u in (urlset.find_all("url", recursive=False) if urlset is not None else []):
                loc = u.find("loc", recursive=False)   # skips <image:loc>, <video:loc>...
                if loc and loc.text.strip():
                    urls.append(loc.text.strip())
        if top:
            urls = list(dict.fromkeys(urls))
            self._status(f"Sitemap: {len(urls)} URLs found")
        return urls

    # ---------- advanced extraction ----------
    def _extract_structured(self, soup):
        types = []
        for s in soup.find_all("script", attrs={"type": re.compile("ld\\+json", re.I)}):
            raw = s.string or s.get_text() or ""
            try:
                data = json.loads(raw)
            except Exception:
                continue
            for obj in (data if isinstance(data, list) else [data]):
                if not isinstance(obj, dict):
                    continue
                graph = obj.get("@graph", [obj])
                for node in (graph if isinstance(graph, list) else [graph]):
                    if isinstance(node, dict) and node.get("@type"):
                        t = node["@type"]
                        types.extend(t if isinstance(t, list) else [t])
        if soup.find(attrs={"itemscope": True}):
            types.append("Microdata")
        # dedupe, preserving order
        seen, out = set(), []
        for t in types:
            if t not in seen:
                seen.add(t); out.append(t)
        return out

    def _apply_extractors(self, soup, html, r):
        if not self.extractors:
            return
        tree = None
        for i, ex in enumerate(self.extractors):
            if i >= MAX_CUSTOM:
                break
            typ = ex.get("type", "css").lower()
            expr = ex.get("expr", "")
            val = ""
            if not expr:
                setattr(r, f"custom_{i+1}", ""); continue
            try:
                if typ == "css":
                    nodes = soup.select(expr)
                    val = " | ".join(n.get_text(strip=True) for n in nodes[:5])
                elif typ == "xpath":
                    if tree is None:
                        from lxml import html as lxml_html
                        tree = lxml_html.fromstring(html)
                    nodes = tree.xpath(expr)
                    parts = []
                    for n in nodes[:5]:
                        parts.append(n.text_content().strip() if hasattr(n, "text_content") else str(n).strip())
                    val = " | ".join(parts)
                elif typ == "regex":
                    matches = re.findall(expr, html)
                    flat = [" ".join(m) if isinstance(m, tuple) else m for m in matches[:5]]
                    val = " | ".join(flat)
            except Exception as e:
                val = f"[error: {str(e)[:40]}]"
            setattr(r, f"custom_{i+1}", val)

    # ---------- parsing ----------
    def _parse(self, url, resp, depth):
        r = PageResult(url=url, crawl_depth=depth)
        r.status_code = resp.status_code
        r.content_type = resp.headers.get("Content-Type", "").split(";")[0].strip()
        r.response_ms = int(resp.elapsed.total_seconds() * 1000)
        r.size_kb = round(len(resp.content) / 1024, 1)
        r.x_robots = resp.headers.get("X-Robots-Tag", "").strip()
        r.is_https = "Yes" if url.lower().startswith("https://") else "No"

        if resp.history:
            r.redirect_to = resp.url
            r.redirect_chain = " -> ".join(
                [str(h.status_code) for h in resp.history] + [str(resp.status_code)])
            r.status_code = resp.history[0].status_code

        if "html" not in r.content_type.lower():
            r.indexable, r.indexability_status = "Non-Indexable", "No HTML"
            return r, [], []

        # Render JS if enabled (on the final URL)
        html = resp.text
        if self.render_js and self.renderer:
            try:
                html, vitals = self.renderer.render(resp.url or url,
                                                    timeout=self.timeout + 10, measure=True)
                r.rendered = "Yes"
                r.ps_fcp = vitals.get("fcp", "")
                r.ps_lcp = vitals.get("lcp", "")
                r.ps_cls = vitals.get("cls", "")
                r.ps_ttfb = vitals.get("ttfb", "")
                r.ps_load = vitals.get("load", "")
            except Exception as e:
                r.rendered = f"No ({str(e)[:30]})"

        soup = BeautifulSoup(html, "lxml")
        html_len = max(len(html), 1)

        if soup.html and soup.html.get("lang"):
            r.lang = soup.html.get("lang").strip()

        if soup.title and soup.title.string:
            r.title = " ".join(soup.title.string.split())
            r.title_len = len(r.title)

        for m in soup.find_all("meta"):
            name = (m.get("name") or "").lower()
            content = (m.get("content") or "").strip()
            if name == "description":
                r.meta_description = " ".join(content.split()); r.meta_desc_len = len(r.meta_description)
            elif name == "keywords":
                r.meta_keywords = content
            elif name == "robots":
                r.meta_robots = content
            elif name == "viewport":
                r.viewport = content
            if m.get("charset"):
                r.charset = m.get("charset").strip()
        if not r.charset and resp.encoding:
            r.charset = resp.encoding

        h1s = [h.get_text(strip=True) for h in soup.find_all("h1")]
        h2s = [h.get_text(strip=True) for h in soup.find_all("h2")]
        r.h1 = " | ".join(h1s[:3]); r.h1_count = len(h1s)
        r.h2 = " | ".join(h2s[:3]); r.h2_count = len(h2s)

        link_canon = soup.find("link", rel=lambda v: v and "canonical" in v)
        if link_canon and link_canon.get("href"):
            r.canonical = urljoin(url, link_canon["href"])
            r.canonical_status = ("Self" if self._normalize(r.canonical) == self._normalize(url)
                                  else "Canonicalized")
        else:
            r.canonical_status = "Missing"

        rn = soup.find("link", rel=lambda v: v and "next" in v)
        rp = soup.find("link", rel=lambda v: v and "prev" in v)
        if rn and rn.get("href"):
            r.rel_next = urljoin(url, rn["href"])
        if rp and rp.get("href"):
            r.rel_prev = urljoin(url, rp["href"])

        hreflangs = [link.get("hreflang") for link in
                     soup.find_all("link", rel=lambda v: v and "alternate" in v)
                     if link.get("hreflang")]
        r.hreflang = ", ".join(hreflangs); r.hreflang_count = len(hreflangs)

        # structured data (before stripping scripts)
        sd = self._extract_structured(soup)
        r.structured_data = ", ".join(sd); r.structured_data_count = len(sd)

        # custom extraction (on the full HTML)
        self._apply_extractors(soup, html, r)

        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        text = soup.get_text(separator=" ", strip=True)
        r.word_count = len(text.split())
        r.text_ratio = round(len(text) / html_len * 100, 1)
        norm = re.sub(r"\s+", " ", text.lower()).strip()
        r.content_hash = hashlib.md5(norm.encode("utf-8", "ignore")).hexdigest()[:16]

        imgs = soup.find_all("img")
        r.images = len(imgs)
        r.images_no_alt = sum(1 for i in imgs if not (i.get("alt") or "").strip())

        robots_all = (r.meta_robots + " " + r.x_robots).lower()
        if "noindex" in robots_all:
            r.indexable, r.indexability_status = "Non-Indexable", "noindex"
        elif r.canonical_status == "Canonicalized":
            r.indexable, r.indexability_status = "Non-Indexable", "Canonicalized"
        elif r.status_code >= 400:
            r.indexable, r.indexability_status = "Non-Indexable", f"Error ({r.status_code})"
        elif r.redirect_to:
            r.indexable, r.indexability_status = "Non-Indexable", "Redirected"
        else:
            r.indexable, r.indexability_status = "Indexable", ""

        found_links, resource_links = [], []
        internal = external = mixed = 0
        page_https = url.lower().startswith("https://")

        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
                continue
            absolute = self._normalize(urljoin(url, href))
            if not absolute.startswith(("http://", "https://")):
                continue
            if page_https and absolute.startswith("http://"):
                mixed += 1
            if self._same_domain(absolute):
                internal += 1; found_links.append(absolute)
            else:
                external += 1
                if self.include_external_status:
                    found_links.append(("EXT", absolute))

        if self.check_resources:
            for link in soup.find_all("link", rel=lambda v: v and "stylesheet" in v):
                if link.get("href"):
                    u = urljoin(url, link["href"]); resource_links.append(("CSS", u))
                    if page_https and u.startswith("http://"):
                        mixed += 1
            for s in soup.find_all("script", src=True):
                u = urljoin(url, s["src"]); resource_links.append(("JavaScript", u))
                if page_https and u.startswith("http://"):
                    mixed += 1
            for i in imgs:
                if i.get("src"):
                    u = urljoin(url, i["src"]); resource_links.append(("Image", u))
                    if page_https and u.startswith("http://"):
                        mixed += 1

        r.internal_links = internal; r.external_links = external
        r.outlinks = internal + external; r.mixed_content = mixed
        return r, found_links, resource_links

    # ---------- fetch ----------
    def _fetch(self, url, depth, ua):
        try:
            resp = self.session.get(url, timeout=self.timeout, allow_redirects=True)
            return self._parse(url, resp, depth)
        except requests.exceptions.RequestException as e:
            r = PageResult(url=url, crawl_depth=depth, error=str(e)[:200],
                           indexable="Non-Indexable", indexability_status="Error")
            return r, [], []

    def _check_url_status(self, url):
        t = min(self.timeout, 10)   # short timeout: dead links fail faster
        try:
            resp = self.session.head(url, timeout=t, allow_redirects=True)
            code = resp.status_code
            if code >= 400 or code == 405:
                resp = self.session.get(url, timeout=t, allow_redirects=True, stream=True)
                code = resp.status_code; resp.close()
            return code
        except requests.exceptions.RequestException:
            return 0

    def _check_external(self, url):
        with self.lock:
            if url in self.external_checked:
                return
        code = self._check_url_status(url)
        with self.lock:
            self.external_checked[url] = code

    def _check_resource(self, url):
        rtype, source = self._resource_urls[url]
        code = self._check_url_status(url)
        with self.lock:
            self.resources[url] = ResourceResult(url=url, type=rtype, status_code=code, source=source)

    # ---------- crawl ----------
    def crawl(self):
        ua = self.session.headers["User-Agent"]
        # seeds depending on mode
        if self.mode == "sitemap":
            seeds = self.load_sitemap_urls(self.start_url)
            if not seeds:
                raise RuntimeError(f"No URLs found in the sitemap for {self.start_url}. "
                                   "Check the URL (e.g. https://example.com/sitemap.xml).")
        elif self.mode == "list":
            seeds = self.seed_urls
        else:
            seeds = [self.start_url]
        for s in seeds:
            # in sitemap/list mode keep the exact URL (e.g. WordPress trailing slash)
            n = self._normalize(s) if self.mode == "spider" else urldefrag(s)[0]
            if n not in self.seen:
                self.seen.add(n); self.queue.append((n, 0))

        external_queue = []
        processed = 0
        try:
            from concurrent.futures import wait, FIRST_COMPLETED
            ex = ThreadPoolExecutor(max_workers=self.threads)
            inflight = {}   # future -> (url, depth)

            def submit_more():
                while (self.queue and len(inflight) < self.threads
                       and (processed + len(inflight)) < self.max_urls
                       and not self.stop_event.is_set()):
                    u, d = self.queue.popleft()
                    inflight[ex.submit(self._fetch, u, d, ua)] = (u, d)

            submit_more()
            while inflight and not self.stop_event.is_set():
                done, _ = wait(list(inflight), return_when=FIRST_COMPLETED)
                for fut in done:
                    inflight.pop(fut, None)
                    result, links, resource_links = fut.result()
                    processed += 1
                    self.results[result.url] = result
                    # emit each URL immediately (live view)
                    if self.on_result:
                        self.on_result(result)
                    else:
                        print(f"[{processed:>5}] {result.status_code or 'ERR':>3} "
                              f"d{result.crawl_depth} {result.url[:88]}")

                    for rtype, ru in resource_links:
                        if ru not in self._resource_urls:
                            self._resource_urls[ru] = (rtype, result.url)

                    for link in links:
                        if isinstance(link, tuple):
                            external_queue.append(link[1]); continue
                        self.inlink_counter[link] = self.inlink_counter.get(link, 0) + 1
                        if (self.follow_links and link not in self.seen
                                and len(self.seen) < self.max_urls
                                and self._passes_filters(link) and self._allowed(link, ua)):
                            self.seen.add(link); self.queue.append((link, result.crawl_depth + 1))

                if self.on_progress:
                    self.on_progress(processed, len(self.queue) + len(inflight))
                if self.delay:
                    time.sleep(self.delay)
                submit_more()

            ex.shutdown(wait=False, cancel_futures=True)

            for url, r in self.results.items():
                r.inlinks = self.inlink_counter.get(url, 0) or self.inlink_counter.get(self._normalize(url), 0)

            if self.include_external_status and external_queue and not self.stop_event.is_set():
                uniq = list(set(external_queue))
                if self.on_status:
                    self.on_status(f"Checking {len(uniq)} external links...")
                with ThreadPoolExecutor(max_workers=min(max(self.threads, 10), 20)) as ex:
                    ex.map(self._check_external, uniq)

            if self.check_resources and self._resource_urls and not self.stop_event.is_set():
                if self.on_status:
                    self.on_status(f"Checking {len(self._resource_urls)} resources...")
                with ThreadPoolExecutor(max_workers=min(max(self.threads, 10), 20)) as ex:
                    ex.map(self._check_resource, list(self._resource_urls.keys()))
        finally:
            if self.renderer:
                self.renderer.quit_all()
        return self.results

    # ---------- PageSpeed / Core Web Vitals ----------
    def fetch_pagespeed(self, api_key, strategy="mobile", limit=10, on_each=None):
        """Fill ps_* on HTML 200 pages (up to 'limit'). Requires api_key."""
        if not api_key:
            return 0
        targets = [r for r in self.results.values()
                   if 200 <= r.status_code < 300 and "html" in (r.content_type or "").lower()][:limit]
        endpoint = "https://www.googleapis.com/pagespeedonline/v5/runPagespeed"
        done = 0
        for r in targets:
            if self.stop_event.is_set():
                break
            try:
                resp = self.session.get(endpoint, params={
                    "url": r.url, "key": api_key, "strategy": strategy,
                    "category": "performance"}, timeout=60)
                data = resp.json()
                lh = data.get("lighthouseResult", {})
                cats = lh.get("categories", {}).get("performance", {})
                audits = lh.get("audits", {})
                if cats.get("score") is not None:
                    r.ps_performance = str(int(cats["score"] * 100))
                r.ps_lcp = audits.get("largest-contentful-paint", {}).get("displayValue", "")
                r.ps_cls = audits.get("cumulative-layout-shift", {}).get("displayValue", "")
                r.ps_fcp = audits.get("first-contentful-paint", {}).get("displayValue", "")
            except Exception as e:
                r.ps_performance = f"err:{str(e)[:20]}"
            done += 1
            if on_each:
                on_each(done, len(targets), r)
        return done

    # ---------- analysis ----------
    def duplicates(self):
        def group(attr):
            buckets = {}
            for r in self.results.values():
                if "html" not in (r.content_type or "").lower():
                    continue
                val = getattr(r, attr)
                if val:
                    buckets.setdefault(val, []).append(r.url)
            return {v: urls for v, urls in buckets.items() if len(urls) > 1}
        return {
            "titles": group("title"),
            "descriptions": group("meta_description"),
            "h1": group("h1"),
            "content": group("content_hash"),
        }

    def summary_counts(self):
        res = self.results.values()
        html = [r for r in res if "html" in (r.content_type or "").lower()]
        codes = {"2xx": 0, "3xx": 0, "4xx": 0, "5xx": 0, "err": 0}
        for r in res:
            c = r.status_code
            if c == 0: codes["err"] += 1
            elif 200 <= c < 300: codes["2xx"] += 1
            elif 300 <= c < 400: codes["3xx"] += 1
            elif 400 <= c < 500: codes["4xx"] += 1
            elif c >= 500: codes["5xx"] += 1
        dups = self.duplicates()
        issues = {
            "Missing Title": sum(1 for r in html if not r.title),
            "Title > 60": sum(1 for r in html if r.title_len > 60),
            "Duplicate Titles": sum(len(v) for v in dups["titles"].values()),
            "Missing Meta Description": sum(1 for r in html if not r.meta_description),
            "Meta Desc > 160": sum(1 for r in html if r.meta_desc_len > 160),
            "Duplicate Descriptions": sum(len(v) for v in dups["descriptions"].values()),
            "Missing H1": sum(1 for r in html if r.h1_count == 0),
            "Multiple H1": sum(1 for r in html if r.h1_count > 1),
            "Duplicate Content": sum(len(v) for v in dups["content"].values()),
            "Non-Indexable": sum(1 for r in html if r.indexable == "Non-Indexable"),
            "Canonicalized": sum(1 for r in html if r.canonical_status == "Canonicalized"),
            "Images Missing Alt": sum(1 for r in html if r.images_no_alt > 0),
            "Mixed Content": sum(1 for r in html if r.mixed_content > 0),
            "With Structured Data": sum(1 for r in html if r.structured_data_count > 0),
        }
        broken_res = sum(1 for rr in self.resources.values() if rr.status_code == 0 or rr.status_code >= 400)
        broken_ext = sum(1 for c in self.external_checked.values() if c == 0 or c >= 400)
        return {"total": len(self.results), "html": len(html), "codes": codes, "issues": issues,
                "broken_resources": broken_res, "broken_external": broken_ext,
                "resources": len(self.resources), "external": len(self.external_checked)}

    # ---------- export ----------
    def export_csv(self, path):
        fields = list(PageResult().__dict__.keys())
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
            for r in self.results.values():
                w.writerow(asdict(r))

    def export_xlsx(self, path):
        """Export to Excel with several sheets (Overview, Internal, Resources, External, Duplicates)."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill("solid", fgColor="1F9E6E")

        ws = wb.active; ws.title = "Overview"
        s = self.summary_counts()
        ws.append(["SilentRat SEO Spider", self.domain]); ws.append([])
        ws.append(["URLs", s["total"]]); ws.append(["HTML", s["html"]])
        for k in ("2xx", "3xx", "4xx", "5xx", "err"):
            ws.append([k, s["codes"][k]])
        ws.append([]); ws.append(["SEO Issues", ""])
        for k, v in s["issues"].items():
            ws.append([k, v])

        # Internal: every PageResult column
        fields = list(PageResult().__dict__.keys())
        wi = wb.create_sheet("Internal"); wi.append(fields)
        for cell in wi[1]:
            cell.font = hdr_font; cell.fill = hdr_fill
        for r in self.results.values():
            wi.append([str(getattr(r, f, "")) for f in fields])

        wr = wb.create_sheet("Resources"); wr.append(["Address", "Type", "Status", "Found On"])
        for cell in wr[1]:
            cell.font = hdr_font; cell.fill = hdr_fill
        for rr in self.resources.values():
            wr.append([rr.url, rr.type, rr.status_code, rr.source])

        we = wb.create_sheet("External"); we.append(["Address", "Status"])
        for cell in we[1]:
            cell.font = hdr_font; cell.fill = hdr_fill
        for u, c in self.external_checked.items():
            we.append([u, c])

        wd = wb.create_sheet("Duplicates"); wd.append(["Type", "Value/Hash", "Count", "URLs"])
        for cell in wd[1]:
            cell.font = hdr_font; cell.fill = hdr_fill
        labels = {"titles": "Title", "descriptions": "Meta Description", "h1": "H1", "content": "Content"}
        for key, groups in self.duplicates().items():
            for value, urls in groups.items():
                wd.append([labels[key], (value or "")[:200], len(urls), "  |  ".join(urls)])

        wb.save(path)

    def generate_sitemap(self, path):
        urls = [r.url for r in self.results.values()
                if r.indexable == "Indexable" and 200 <= r.status_code < 300
                and "html" in (r.content_type or "").lower()]
        lines = ['<?xml version="1.0" encoding="UTF-8"?>',
                 '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">']
        for u in urls:
            lines.append(f"  <url><loc>{xml_escape(u)}</loc></url>")
        lines.append("</urlset>")
        with open(path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return len(urls)


# ---------- CLI ----------
def main():
    p = argparse.ArgumentParser(description="SilentRat SEO Spider (advanced CLI)")
    p.add_argument("url", help="Start URL (or sitemap URL in sitemap mode)")
    p.add_argument("--max", type=int, default=1_000_000)
    p.add_argument("--delay", type=float, default=0.0)
    p.add_argument("--threads", type=int, default=5)
    p.add_argument("--out", default="silentrat_report.csv")
    p.add_argument("--xlsx", default="", help="Write an Excel report to this path")
    p.add_argument("--sitemap-out", default="", help="Write a sitemap.xml to this path")
    p.add_argument("--mode", choices=["spider", "list", "sitemap"], default="spider")
    p.add_argument("--list-file", default="", help="File of URLs (list mode)")
    p.add_argument("--render", action="store_true", help="Render JavaScript (headless)")
    p.add_argument("--no-robots", action="store_true")
    p.add_argument("--no-resources", action="store_true")
    p.add_argument("--include", default="")
    p.add_argument("--exclude", default="")
    p.add_argument("--psi-key", default="", help="PageSpeed Insights API key")
    args = p.parse_args()

    if not args.url.startswith(("http://", "https://")):
        args.url = "https://" + args.url
    seeds = []
    if args.mode == "list" and args.list_file:
        with open(args.list_file, encoding="utf-8") as f:
            seeds = [ln.strip() for ln in f if ln.strip()]

    spider = SeoSpider(
        start_url=args.url, max_urls=args.max, delay=args.delay, threads=args.threads,
        respect_robots=not args.no_robots, check_resources=not args.no_resources,
        include_pattern=args.include, exclude_pattern=args.exclude,
        mode=args.mode, seed_urls=seeds, render_js=args.render)
    t0 = time.time(); spider.crawl(); elapsed = time.time() - t0

    if args.psi_key:
        print("[pagespeed] Fetching Core Web Vitals...")
        spider.fetch_pagespeed(args.psi_key)

    s = spider.summary_counts()
    print("\n" + "=" * 55)
    print(f"  SilentRat - {spider.domain}  (mode {args.mode}"
          f"{', JS' if args.render else ''})")
    print("=" * 55)
    print(f"  URLs: {s['total']}  HTML: {s['html']}")
    print(f"  Status: 2xx={s['codes']['2xx']} 3xx={s['codes']['3xx']} "
          f"4xx={s['codes']['4xx']} 5xx={s['codes']['5xx']} err={s['codes']['err']}")
    print(f"  Resources: {s['resources']} (broken {s['broken_resources']})  "
          f"External: {s['external']} (broken {s['broken_external']})")
    for k, v in s["issues"].items():
        if v:
            print(f"    {k}: {v}")
    print("=" * 55)
    spider.export_csv(args.out); print(f"[export] {args.out}")
    if args.xlsx:
        spider.export_xlsx(args.xlsx); print(f"[export] {args.xlsx}")
    if args.sitemap_out:
        n = spider.generate_sitemap(args.sitemap_out); print(f"[sitemap] {n} URLs -> {args.sitemap_out}")
    print(f"[done] {s['total']} URLs in {elapsed:.1f}s")


if __name__ == "__main__":
    main()
