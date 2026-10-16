#!/usr/bin/env python3
"""
SilentRat - desktop SEO Spider (full GUI + advanced features).

Standalone SEO crawler in the style of Screaming Frog, with no URL limit.
Engine: seo_spider.py

Run:  python silentrat.py   (or SilentRat.bat / SilentRat.exe)
"""

import os
import queue
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

from seo_spider import SeoSpider, PageResult, DEFAULT_UA, MAX_CUSTOM

ACCENT = "#1f9e6e"
ACCENT_DK = "#17704f"
BG = "#f4f5f7"
HEADER_BG = "#2b2f36"


def _is_html(r):
    return "html" in (r.content_type or "").lower()


# column mode: live (PageResult) | external | resources | duplicates
TABS = {
    "Internal": dict(mode="live", filter=lambda r: True, cols=[
        ("url", "Address", 380), ("status_code", "Status", 55),
        ("content_type", "Type", 80), ("indexable", "Indexability", 105),
        ("indexability_status", "Reason", 120), ("title", "Title 1", 220),
        ("word_count", "Words", 55), ("size_kb", "Size KB", 60),
        ("rendered", "JS", 50), ("inlinks", "Inlinks", 55),
        ("outlinks", "Outlinks", 60), ("crawl_depth", "Depth", 50)]),
    "Response Codes": dict(mode="live", filter=lambda r: True, cols=[
        ("url", "Address", 430), ("status_code", "Status", 70),
        ("content_type", "Type", 100), ("redirect_to", "Redirect URL", 320),
        ("redirect_chain", "Chain", 110), ("response_ms", "Time ms", 70)]),
    "Page Titles": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 430), ("title", "Title 1", 440), ("title_len", "Length", 65)]),
    "Meta Description": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 400), ("meta_description", "Meta Description 1", 450),
        ("meta_desc_len", "Length", 65)]),
    "Meta Keywords": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 460), ("meta_keywords", "Meta Keywords 1", 420)]),
    "H1": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 450), ("h1", "H1-1", 430), ("h1_count", "H1 Count", 75)]),
    "H2": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 450), ("h2", "H2-1", 430), ("h2_count", "H2 Count", 75)]),
    "Images": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 520), ("images", "Images", 80), ("images_no_alt", "Missing Alt", 110)]),
    "Canonicals": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 420), ("canonical", "Canonical", 420), ("canonical_status", "Status", 120)]),
    "Directives": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 390), ("indexable", "Indexability", 105),
        ("indexability_status", "Reason", 140), ("meta_robots", "Meta Robots", 170),
        ("x_robots", "X-Robots-Tag", 140)]),
    "Hreflang": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 440), ("hreflang", "hreflang", 350),
        ("hreflang_count", "Count", 60), ("lang", "HTML lang", 85)]),
    "Structured Data": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 430), ("structured_data", "Types (JSON-LD / Microdata)", 430),
        ("structured_data_count", "Count", 70)]),
    "Custom Extraction": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 360), ("custom_1", "Custom 1", 150), ("custom_2", "Custom 2", 150),
        ("custom_3", "Custom 3", 150), ("custom_4", "Custom 4", 150), ("custom_5", "Custom 5", 150)]),
    "Web Vitals": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 400), ("ps_fcp", "FCP", 95), ("ps_lcp", "LCP", 95),
        ("ps_cls", "CLS", 80), ("ps_ttfb", "TTFB", 95), ("ps_load", "Load", 95),
        ("rendered", "Rendered", 90)]),
    "Security": dict(mode="live", filter=_is_html, cols=[
        ("url", "Address", 540), ("is_https", "HTTPS", 70), ("mixed_content", "Mixed Content", 120)]),
    "Resources": dict(mode="resources", filter=None, cols=[
        ("url", "Address", 520), ("type", "Type", 110), ("status_code", "Status", 80),
        ("source", "Found On", 360)]),
    "External": dict(mode="external", filter=None, cols=[
        ("url", "Address", 640), ("status_code", "Status", 100)]),
    "Duplicates": dict(mode="duplicates", filter=None, cols=[
        ("dtype", "Type", 150), ("value", "Value / Hash", 360),
        ("count", "Count", 70), ("urls", "URLs", 560)]),
}

NUM_COLS = {"status_code", "title_len", "meta_desc_len", "h1_count", "h2_count",
            "word_count", "size_kb", "inlinks", "outlinks", "crawl_depth", "images",
            "images_no_alt", "hreflang_count", "mixed_content", "response_ms",
            "structured_data_count", "count"}


class SilentRatApp:
    def __init__(self, root):
        self.root = root
        root.title("SilentRat SEO Spider")
        root.geometry("1320x780")
        root.configure(bg=BG)
        root.minsize(1000, 620)

        self.spider = None
        self.result_queue = queue.Queue()
        self.trees = {}
        self.tab_rows = {}
        self.running = False
        self.list_urls = []
        # Features enabled by default. JS rendering is OFF for performance
        # (it opens a browser per URL): enable it in Config only for JS-heavy sites.
        # Web Vitals are measured locally during rendering, with no PageSpeed API.
        self.config = dict(
            user_agent=DEFAULT_UA, delay=0.0, include="", exclude="",
            check_resources=True, check_external=True,
            render_js=False, render_wait=2.0,
            extractors=[dict(name="", type="css", expr="") for _ in range(MAX_CUSTOM)],
        )

        self._setup_style()
        self._build_header()
        self._build_toolbar()
        self._build_tabs()
        self._build_statusbar()
        self.root.after(120, self._drain_queue)

    def _setup_style(self):
        st = ttk.Style()
        try:
            st.theme_use("clam")
        except tk.TclError:
            pass
        st.configure("Treeview", rowheight=22, font=("Segoe UI", 9),
                     background="white", fieldbackground="white")
        st.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"),
                     background="#e4e6ea", relief="flat")
        st.map("Treeview.Heading", background=[("active", "#d0d3d8")])
        st.configure("TNotebook", background=BG, borderwidth=0)
        st.configure("TNotebook.Tab", padding=(9, 6), font=("Segoe UI", 9))
        st.map("TNotebook.Tab", background=[("selected", ACCENT), ("!selected", "#d7dae0")],
               foreground=[("selected", "white"), ("!selected", "#2b2f36")])
        st.configure("Accent.TButton", font=("Segoe UI", 9, "bold"), foreground="white",
                     background=ACCENT, borderwidth=0, padding=(16, 6))
        st.map("Accent.TButton", background=[("active", ACCENT_DK), ("disabled", "#9bbfb0")])
        st.configure("Stop.TButton", font=("Segoe UI", 9, "bold"), foreground="white",
                     background="#c0392b", borderwidth=0, padding=(16, 6))
        st.map("Stop.TButton", background=[("active", "#962d22")])

    def _draw_logo(self, c):
        c.create_oval(10, 16, 40, 40, fill="#5b6470", outline="")
        c.create_oval(30, 12, 50, 32, fill="#6b7482", outline="")
        c.create_oval(30, 6, 42, 18, fill="#7d8695", outline="")
        c.create_oval(40, 6, 52, 18, fill="#7d8695", outline="")
        c.create_oval(33, 9, 39, 15, fill="#c98da6", outline="")
        c.create_oval(43, 9, 49, 15, fill="#c98da6", outline="")
        c.create_oval(44, 19, 48, 23, fill="#1b1f24", outline="")
        c.create_oval(49, 22, 53, 26, fill="#c0392b", outline="")
        c.create_line(50, 24, 60, 20, fill="#2b2f36")
        c.create_line(50, 25, 60, 25, fill="#2b2f36")
        c.create_line(50, 26, 60, 30, fill="#2b2f36")
        c.create_line(12, 34, 2, 46, fill="#5b6470", width=3, smooth=True)

    def _build_header(self):
        h = tk.Frame(self.root, bg=HEADER_BG, height=58); h.pack(fill="x"); h.pack_propagate(False)
        logo = tk.Canvas(h, width=64, height=52, bg=HEADER_BG, highlightthickness=0)
        logo.pack(side="left", padx=(14, 6), pady=3); self._draw_logo(logo)
        tk.Label(h, text="SilentRat", bg=HEADER_BG, fg="white",
                 font=("Segoe UI", 17, "bold")).pack(side="left", pady=(10, 0))
        tk.Label(h, text="SEO Spider", bg=HEADER_BG, fg=ACCENT,
                 font=("Segoe UI", 11, "bold")).pack(side="left", padx=(6, 0), pady=(14, 0))
        tk.Label(h, text="standalone · unlimited · JS · Web Vitals", bg=HEADER_BG,
                 fg="#8a909a", font=("Segoe UI", 8)).pack(side="right", padx=14, pady=(20, 0))

    def _build_toolbar(self):
        bar = tk.Frame(self.root, bg="#e9ebee", height=46); bar.pack(fill="x"); bar.pack_propagate(False)
        tk.Label(bar, text="Mode:", bg="#e9ebee", font=("Segoe UI", 9)).pack(side="left", padx=(10, 2))
        self.mode_var = tk.StringVar(value="Spider")
        mode_cb = ttk.Combobox(bar, textvariable=self.mode_var, width=8, state="readonly",
                               values=["Spider", "List", "Sitemap"])
        mode_cb.pack(side="left"); mode_cb.bind("<<ComboboxSelected>>", lambda e: self._on_mode())
        tk.Label(bar, text="URL:", bg="#e9ebee", font=("Segoe UI", 9)).pack(side="left", padx=(8, 4))
        self.url_var = tk.StringVar(value="https://")
        self.url_entry = tk.Entry(bar, textvariable=self.url_var, font=("Segoe UI", 10), width=36)
        self.url_entry.pack(side="left", padx=4, pady=8, ipady=3)
        self.url_entry.bind("<Return>", lambda e: self.start())
        self.list_btn = ttk.Button(bar, text="URLs...", command=self.edit_list, width=7)
        tk.Label(bar, text="Max:", bg="#e9ebee", font=("Segoe UI", 9)).pack(side="left", padx=(8, 2))
        self.max_var = tk.StringVar(value="0")
        tk.Entry(bar, textvariable=self.max_var, width=7, font=("Segoe UI", 9)).pack(side="left")
        tk.Label(bar, text="Threads:", bg="#e9ebee", font=("Segoe UI", 9)).pack(side="left", padx=(8, 2))
        self.threads_var = tk.StringVar(value="5")
        tk.Entry(bar, textvariable=self.threads_var, width=4, font=("Segoe UI", 9)).pack(side="left")
        self.robots_var = tk.BooleanVar(value=True)
        tk.Checkbutton(bar, text="robots", variable=self.robots_var, bg="#e9ebee",
                       font=("Segoe UI", 8)).pack(side="left", padx=(6, 0))
        self.start_btn = ttk.Button(bar, text="Start", style="Accent.TButton", command=self.start)
        self.start_btn.pack(side="left", padx=(10, 3), pady=6)
        self.stop_btn = ttk.Button(bar, text="Stop", style="Stop.TButton", command=self.stop, state="disabled")
        self.stop_btn.pack(side="left", padx=2)
        ttk.Button(bar, text="Config", command=self.open_config).pack(side="left", padx=2)
        ttk.Button(bar, text="Export", command=self.export_menu).pack(side="left", padx=2)
        ttk.Button(bar, text="Sitemap", command=self.gen_sitemap).pack(side="left", padx=2)
        ttk.Button(bar, text="Clear", command=self.clear).pack(side="left", padx=2)

    def _on_mode(self):
        mode = self.mode_var.get()
        if mode == "List":
            self.list_btn.pack(side="left", padx=2, after=self.url_entry)
            self.url_entry.config(state="disabled")
        else:
            self.list_btn.pack_forget()
            self.url_entry.config(state="normal")
            if mode == "Sitemap":
                self.status_var.set("Sitemap mode: enter the sitemap.xml URL or just the domain "
                                    "(auto-detected via robots.txt).")

    def edit_list(self):
        d = tk.Toplevel(self.root); d.title("URL List"); d.configure(bg=BG)
        d.transient(self.root); d.grab_set()
        tk.Label(d, text="One URL per line:", bg=BG, font=("Segoe UI", 9)).pack(anchor="w", padx=10, pady=(10, 2))
        txt = scrolledtext.ScrolledText(d, width=70, height=18, font=("Consolas", 9))
        txt.pack(padx=10, pady=4); txt.insert("1.0", "\n".join(self.list_urls))

        def save():
            self.list_urls = [l.strip() for l in txt.get("1.0", "end").splitlines() if l.strip()]
            d.destroy()
        ttk.Button(d, text="Save", style="Accent.TButton", command=save).pack(pady=8)

    def _build_tabs(self):
        cont = tk.Frame(self.root, bg=BG); cont.pack(fill="both", expand=True, padx=8, pady=(6, 4))
        filt = tk.Frame(cont, bg=BG); filt.pack(fill="x")
        tk.Label(filt, text="Filter:", bg=BG, font=("Segoe UI", 9)).pack(side="left")
        self.filter_var = tk.StringVar()
        fe = tk.Entry(filt, textvariable=self.filter_var, font=("Segoe UI", 9), width=44)
        fe.pack(side="left", padx=6, pady=4, ipady=2); fe.bind("<KeyRelease>", lambda e: self._apply_filter())
        self.row_info = tk.Label(filt, text="", bg=BG, fg="#888", font=("Segoe UI", 8)); self.row_info.pack(side="left")

        self.nb = ttk.Notebook(cont); self.nb.pack(fill="both", expand=True, pady=(4, 0))
        self.nb.bind("<<NotebookTabChanged>>", lambda e: self._apply_filter())

        self.ov_frame = tk.Frame(self.nb, bg="white"); self.nb.add(self.ov_frame, text="Overview")
        self._build_overview(self.ov_frame)

        for name, cfg in TABS.items():
            frame = tk.Frame(self.nb, bg="white"); self.nb.add(frame, text=name)
            cols = [c[0] for c in cfg["cols"]]
            tree = ttk.Treeview(frame, columns=cols, show="headings", selectmode="browse")
            for key, head, width in cfg["cols"]:
                tree.heading(key, text=head, command=lambda t=name, k=key: self._sort(t, k))
                tree.column(key, width=width, anchor=("center" if key in NUM_COLS else "w"),
                            stretch=(key in ("url", "urls")))
            vs = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
            hs = ttk.Scrollbar(frame, orient="horizontal", command=tree.xview)
            tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
            tree.grid(row=0, column=0, sticky="nsew"); vs.grid(row=0, column=1, sticky="ns")
            hs.grid(row=1, column=0, sticky="ew")
            frame.rowconfigure(0, weight=1); frame.columnconfigure(0, weight=1)
            for tag, col in (("ok", "#1b1f24"), ("redirect", "#b8860b"),
                             ("error", "#c0392b"), ("noindex", "#8e44ad")):
                tree.tag_configure(tag, foreground=col)
            self.trees[name] = tree; self.tab_rows[name] = []

    def _build_overview(self, parent):
        left = tk.Frame(parent, bg="white"); left.pack(side="left", fill="both", expand=True, padx=12, pady=12)
        tk.Label(left, text="Crawl Summary", bg="white", font=("Segoe UI", 13, "bold"),
                 fg="#2b2f36").pack(anchor="w")
        self.ov_text = tk.Text(left, width=46, height=28, font=("Consolas", 10),
                               bg="#fafbfc", relief="flat", bd=0)
        self.ov_text.pack(anchor="w", pady=(8, 0), fill="both", expand=True)
        self.ov_text.insert("1.0", "Enter a URL and click Start.")
        self.ov_text.config(state="disabled")
        right = tk.Frame(parent, bg="white"); right.pack(side="left", fill="both", expand=True, padx=12, pady=12)
        tk.Label(right, text="Status codes", bg="white", font=("Segoe UI", 13, "bold"),
                 fg="#2b2f36").pack(anchor="w")
        self.ov_canvas = tk.Canvas(right, bg="white", highlightthickness=0, height=330)
        self.ov_canvas.pack(fill="both", expand=True, pady=(8, 0))

    def _build_statusbar(self):
        sb = tk.Frame(self.root, bg="#2b2f36", height=28); sb.pack(fill="x", side="bottom"); sb.pack_propagate(False)
        self.status_var = tk.StringVar(value="Ready. Enter a URL and click Start.")
        tk.Label(sb, textvariable=self.status_var, bg="#2b2f36", fg="white",
                 font=("Segoe UI", 9)).pack(side="left", padx=10)
        self.counts_var = tk.StringVar(value="")
        tk.Label(sb, textvariable=self.counts_var, bg="#2b2f36", fg=ACCENT,
                 font=("Segoe UI", 9, "bold")).pack(side="right", padx=10)
        self.progress = ttk.Progressbar(sb, mode="indeterminate", length=130)

    # ---------- config ----------
    def open_config(self):
        d = tk.Toplevel(self.root); d.title("Advanced Settings"); d.configure(bg=BG)
        d.transient(self.root); d.grab_set(); d.resizable(False, False)
        pad = dict(padx=8, pady=4)
        row = 0
        def label(t, r):
            tk.Label(d, text=t, bg=BG, font=("Segoe UI", 9, "bold")).grid(row=r, column=0, columnspan=3, sticky="w", padx=8, pady=(8, 0))

        tk.Label(d, text="User-Agent:", bg=BG).grid(row=row, column=0, sticky="e", **pad)
        ua = tk.StringVar(value=self.config["user_agent"])
        tk.Entry(d, textvariable=ua, width=56).grid(row=row, column=1, columnspan=2, sticky="w", **pad); row += 1
        tk.Label(d, text="Delay (s):", bg=BG).grid(row=row, column=0, sticky="e", **pad)
        delay = tk.StringVar(value=str(self.config["delay"]))
        tk.Entry(d, textvariable=delay, width=8).grid(row=row, column=1, sticky="w", **pad); row += 1
        tk.Label(d, text="Include (regex):", bg=BG).grid(row=row, column=0, sticky="e", **pad)
        inc = tk.StringVar(value=self.config["include"])
        tk.Entry(d, textvariable=inc, width=44).grid(row=row, column=1, columnspan=2, sticky="w", **pad); row += 1
        tk.Label(d, text="Exclude (regex):", bg=BG).grid(row=row, column=0, sticky="e", **pad)
        exc = tk.StringVar(value=self.config["exclude"])
        tk.Entry(d, textvariable=exc, width=44).grid(row=row, column=1, columnspan=2, sticky="w", **pad); row += 1
        cr = tk.BooleanVar(value=self.config["check_resources"])
        ce = tk.BooleanVar(value=self.config["check_external"])
        tk.Checkbutton(d, text="Check CSS/JS/image resources", variable=cr, bg=BG).grid(row=row, column=1, sticky="w"); row += 1
        tk.Checkbutton(d, text="Check external links", variable=ce, bg=BG).grid(row=row, column=1, sticky="w"); row += 1

        label("JavaScript rendering + local Web Vitals (headless Chrome/Edge)", row); row += 1
        rj = tk.BooleanVar(value=self.config["render_js"])
        tk.Checkbutton(d, text="Enable JS rendering (measures FCP/LCP/CLS/TTFB/Load with no API)",
                       variable=rj, bg=BG).grid(row=row, column=1, columnspan=2, sticky="w")
        tk.Label(d, text="Wait (s):", bg=BG).grid(row=row, column=0, sticky="e", **pad)
        rw = tk.StringVar(value=str(self.config["render_wait"]))
        tk.Entry(d, textvariable=rw, width=6).grid(row=row, column=2, sticky="w"); row += 1

        label("Custom extraction (CSS / XPath / Regex)", row); row += 1
        ex_vars = []
        for i in range(MAX_CUSTOM):
            cfg = self.config["extractors"][i]
            nm = tk.StringVar(value=cfg["name"]); tp = tk.StringVar(value=cfg["type"]); ep = tk.StringVar(value=cfg["expr"])
            tk.Entry(d, textvariable=nm, width=12).grid(row=row, column=0, **pad)
            ttk.Combobox(d, textvariable=tp, width=6, state="readonly",
                         values=["css", "xpath", "regex"]).grid(row=row, column=1, sticky="w", **pad)
            tk.Entry(d, textvariable=ep, width=40).grid(row=row, column=2, sticky="w", **pad)
            ex_vars.append((nm, tp, ep)); row += 1

        def save():
            self.config["user_agent"] = ua.get().strip() or DEFAULT_UA
            try: self.config["delay"] = float(delay.get())
            except ValueError: self.config["delay"] = 0.0
            self.config["include"] = inc.get().strip(); self.config["exclude"] = exc.get().strip()
            self.config["check_resources"] = cr.get(); self.config["check_external"] = ce.get()
            self.config["render_js"] = rj.get()
            try: self.config["render_wait"] = float(rw.get())
            except ValueError: self.config["render_wait"] = 2.0
            self.config["extractors"] = [dict(name=n.get().strip(), type=t.get(), expr=e.get().strip())
                                         for (n, t, e) in ex_vars]
            self._relabel_custom(); d.destroy()
        ttk.Button(d, text="Save", style="Accent.TButton", command=save).grid(
            row=row, column=2, sticky="e", padx=10, pady=12)

    def _relabel_custom(self):
        tree = self.trees["Custom Extraction"]
        for i in range(MAX_CUSTOM):
            name = self.config["extractors"][i]["name"] or f"Custom {i+1}"
            tree.heading(f"custom_{i+1}", text=name)

    # ---------- crawl ----------
    def start(self):
        if self.running:
            return
        mode_map = {"Spider": "spider", "List": "list", "Sitemap": "sitemap"}
        mode = mode_map[self.mode_var.get()]
        url = self.url_var.get().strip()
        seeds = []
        if mode == "list":
            if not self.list_urls:
                messagebox.showwarning("SilentRat", "Add URLs with the 'URLs...' button."); return
            seeds = [u if u.startswith(("http://", "https://")) else "https://" + u for u in self.list_urls]
            url = seeds[0]
        else:
            if not url or url == "https://":
                messagebox.showwarning("SilentRat", "Enter a valid URL."); return
            if not url.startswith(("http://", "https://")):
                url = "https://" + url; self.url_var.set(url)
        try:
            max_urls = int(self.max_var.get()); threads = int(self.threads_var.get())
        except ValueError:
            messagebox.showwarning("SilentRat", "Max and Threads must be numbers."); return
        if max_urls <= 0:
            max_urls = 1_000_000

        extractors = [e for e in self.config["extractors"] if e["expr"]]
        self.clear()
        self.running = True
        self.start_btn.config(state="disabled"); self.stop_btn.config(state="normal")
        self.url_entry.config(state="disabled")
        self.progress.pack(side="right", padx=8); self.progress.start(12)
        extra = " (JS)" if self.config["render_js"] else ""
        self.status_var.set(f"Crawling [{mode}]{extra} {url} ...")
        try:   # switch to the Internal tab to watch URLs live
            self.nb.select(1 + list(TABS).index("Internal"))
        except Exception:
            pass

        self.spider = SeoSpider(
            start_url=url, max_urls=max_urls, threads=threads, delay=self.config["delay"],
            respect_robots=self.robots_var.get(), user_agent=self.config["user_agent"],
            include_external_status=self.config["check_external"],
            check_resources=self.config["check_resources"],
            include_pattern=self.config["include"], exclude_pattern=self.config["exclude"],
            mode=mode, seed_urls=seeds, render_js=self.config["render_js"],
            render_wait=self.config["render_wait"], extractors=extractors)
        self.spider.on_result = lambda r: self.result_queue.put(("row", r))
        self.spider.on_progress = lambda p, q: self.result_queue.put(("prog", (p, q)))
        self.spider.on_status = lambda t: self.result_queue.put(("status", t))
        threading.Thread(target=self._run_crawl, daemon=True).start()

    def _run_crawl(self):
        try:
            self.spider.crawl()
        except Exception as e:
            self.result_queue.put(("error", str(e)))
        finally:
            self.result_queue.put(("done", None))

    def stop(self):
        if self.spider:
            self.spider.stop_event.set(); self.status_var.set("Stopping..."); self.stop_btn.config(state="disabled")

    def clear(self):
        for name, tree in self.trees.items():
            tree.delete(*tree.get_children()); self.tab_rows[name] = []
        self.counts_var.set(""); self.ov_canvas.delete("all"); self._set_overview("Enter a URL and click Start.")

    # ---------- queue -> UI ----------
    def _drain_queue(self):
        try:
            while True:
                kind, payload = self.result_queue.get_nowait()
                if kind == "row": self._add_result(payload)
                elif kind == "prog":
                    pass  # live detail is set by _add_result with the current URL
                elif kind == "status":
                    self.status_var.set(payload)
                elif kind == "error": messagebox.showerror("SilentRat", f"Error: {payload}")
                elif kind == "done": self._finish()
        except queue.Empty:
            pass
        self.root.after(120, self._drain_queue)

    def _row_tag(self, r):
        if r.status_code >= 400 or r.status_code == 0: return "error"
        if 300 <= r.status_code < 400 or getattr(r, "redirect_to", ""): return "redirect"
        if "noindex" in (getattr(r, "meta_robots", "") + getattr(r, "x_robots", "")).lower(): return "noindex"
        return "ok"

    def _add_result(self, r):
        tag = self._row_tag(r)
        current = self._current_tab()
        no_filter = not self.filter_var.get().strip()
        for name, cfg in TABS.items():
            if cfg["mode"] != "live":
                continue
            if cfg["filter"] and not cfg["filter"](r):
                continue
            tree = self.trees[name]
            values = [getattr(r, c[0], "") for c in cfg["cols"]]
            iid = tree.insert("", "end", values=values, tags=(tag,))
            self.tab_rows[name].append((iid, r))
            # auto-scroll the visible tab to follow progress live
            if name == current and no_filter:
                tree.see(iid)
        total = len(self.spider.results) if self.spider else 0
        code = r.status_code or "ERR"
        self.status_var.set(f"Crawling... {total}  [{code}]  {r.url[:70]}")
        self._update_counts()

    def _update_counts(self):
        res = self.spider.results if self.spider else {}
        total = len(res)
        ok = sum(1 for r in res.values() if 200 <= r.status_code < 300)
        redir = sum(1 for r in res.values() if 300 <= r.status_code < 400)
        err = sum(1 for r in res.values() if r.status_code >= 400 or r.status_code == 0)
        self.counts_var.set(f"Total: {total}   2xx: {ok}   3xx: {redir}   4xx/5xx: {err}")

    def _finish(self):
        self.running = False
        self.progress.stop(); self.progress.pack_forget()
        self.start_btn.config(state="normal"); self.stop_btn.config(state="disabled")
        self.url_entry.config(state="normal" if self.mode_var.get() != "List" else "disabled")
        if self.spider and self.spider.external_checked:
            self._fill_tree("External",
                [(u, c or "ERR") for u, c in sorted(self.spider.external_checked.items())])
        if self.spider and self.spider.resources:
            rows = [(rr.url, rr.type, rr.status_code or "ERR", rr.source)
                    for rr in sorted(self.spider.resources.values(), key=lambda x: x.status_code, reverse=True)]
            self._fill_tree("Resources", rows)
        self._fill_duplicates()
        self._update_counts(); self._update_overview()
        n = len(self.spider.results) if self.spider else 0
        self.status_var.set(f"Done. {n} URLs crawled.")

    def _fill_tree(self, name, rows):
        tree = self.trees[name]; tree.delete(*tree.get_children()); self.tab_rows[name] = []
        for row in rows:
            tag = "ok"
            for v in row:
                if v in ("ERR", 0) or (isinstance(v, int) and v >= 400):
                    tag = "error"; break
            iid = tree.insert("", "end", values=row, tags=(tag,))
            self.tab_rows[name].append((iid, PageResult(url=str(row[0])), row))

    def _fill_duplicates(self):
        if not self.spider:
            return
        dups = self.spider.duplicates()
        labels = {"titles": "Title", "descriptions": "Meta Description", "h1": "H1", "content": "Content (hash)"}
        rows = []
        for key, groups in dups.items():
            for value, urls in groups.items():
                rows.append((labels[key], (value or "")[:120], len(urls), "  |  ".join(urls)))
        self._fill_tree("Duplicates", rows)

    # ---------- overview ----------
    def _set_overview(self, text):
        self.ov_text.config(state="normal"); self.ov_text.delete("1.0", "end")
        self.ov_text.insert("1.0", text); self.ov_text.config(state="disabled")

    def _update_overview(self):
        if not self.spider:
            return
        s = self.spider.summary_counts()
        lines = [f"URLs crawled    : {s['total']}", f"  HTML pages    : {s['html']}",
                 f"Resources       : {s['resources']}  (broken: {s['broken_resources']})",
                 f"External links  : {s['external']}  (broken: {s['broken_external']})",
                 "", "STATUS CODES",
                 f"  2xx : {s['codes']['2xx']}", f"  3xx : {s['codes']['3xx']}",
                 f"  4xx : {s['codes']['4xx']}", f"  5xx : {s['codes']['5xx']}",
                 f"  err : {s['codes']['err']}", "", "SEO ISSUES"]
        for k, v in s["issues"].items():
            lines.append(f" {'!' if v else ' '} {k}: {v}")
        self._set_overview("\n".join(lines)); self._draw_chart(s["codes"])

    def _draw_chart(self, codes):
        c = self.ov_canvas; c.delete("all"); c.update_idletasks()
        data = [("2xx", codes["2xx"], ACCENT), ("3xx", codes["3xx"], "#b8860b"),
                ("4xx", codes["4xx"], "#c0392b"), ("5xx", codes["5xx"], "#8e44ad"),
                ("err", codes["err"], "#7f8c8d")]
        mx = max((v for _, v, _ in data), default=0) or 1
        x = 40; bw = 56; gap = 30; base = 280
        for label, val, color in data:
            bh = int((val / mx) * 220)
            c.create_rectangle(x, base - bh, x + bw, base, fill=color, outline="")
            c.create_text(x + bw / 2, base - bh - 12, text=str(val), font=("Segoe UI", 10, "bold"))
            c.create_text(x + bw / 2, base + 14, text=label, font=("Segoe UI", 9))
            x += bw + gap

    # ---------- interaction ----------
    def _current_tab(self):
        try:
            return self.nb.tab(self.nb.select(), "text")
        except tk.TclError:
            return None

    def _apply_filter(self):
        name = self._current_tab()
        if name == "Overview" or name not in self.trees:
            self.row_info.config(text=""); return
        term = self.filter_var.get().lower().strip()
        tree = self.trees[name]; cfg = TABS[name]
        tree.delete(*tree.get_children()); shown = 0
        for entry in self.tab_rows[name]:
            if len(entry) == 3:
                _, pr, values = entry
            else:
                _, r = entry
                values = [getattr(r, c[0], "") for c in cfg["cols"]]
            if term and term not in " ".join(str(v).lower() for v in values):
                continue
            tree.insert("", "end", values=values, tags=(self._tag_for(entry, values),)); shown += 1
        self.row_info.config(text=f"{shown} rows")

    def _tag_for(self, entry, values):
        if len(entry) == 2:
            return self._row_tag(entry[1])
        for v in values:
            if v in ("ERR", 0) or (isinstance(v, int) and v >= 400):
                return "error"
        return "ok"

    def _sort(self, name, key):
        cfg = TABS[name]; rows = self.tab_rows[name]
        rev = getattr(self, "_sort_rev", {}).get((name, key), False)
        idx = [c[0] for c in cfg["cols"]].index(key)

        def sk(entry):
            v = entry[2][idx] if len(entry) == 3 else getattr(entry[1], key, "")
            return (0, v) if isinstance(v, (int, float)) else (1, str(v).lower())
        rows.sort(key=sk, reverse=rev)
        if not hasattr(self, "_sort_rev"):
            self._sort_rev = {}
        self._sort_rev[(name, key)] = not rev
        self._apply_filter()

    # ---------- export ----------
    def export_menu(self):
        if not self.spider or not self.spider.results:
            messagebox.showinfo("SilentRat", "No data yet. Run a crawl first."); return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx"), ("CSV (full report)", "*.csv")],
            initialfile="silentrat_report.xlsx")
        if not path:
            return
        if path.lower().endswith(".csv"):
            self.spider.export_csv(path)
        else:
            self._export_xlsx(path)
        messagebox.showinfo("SilentRat", f"Exported:\n{path}")

    def _export_xlsx(self, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook(); ws0 = wb.active; ws0.title = "Overview"
        s = self.spider.summary_counts()
        ws0.append(["SilentRat SEO Spider", self.spider.domain])
        ws0.append([]); ws0.append(["URLs", s["total"]]); ws0.append(["HTML", s["html"]])
        for k in ("2xx", "3xx", "4xx", "5xx", "err"):
            ws0.append([k, s["codes"][k]])
        ws0.append([]); ws0.append(["SEO Issues", ""])
        for k, v in s["issues"].items():
            ws0.append([k, v])
        hdr_font = Font(bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F9E6E")
        for name, cfg in TABS.items():
            rows = self.tab_rows[name]
            if not rows:
                continue
            ws = wb.create_sheet(name[:31])
            headers = [c[1] for c in cfg["cols"]]
            ws.append(headers)
            for cell in ws[1]:
                cell.font = hdr_font; cell.fill = hdr_fill
            for entry in rows:
                if len(entry) == 3:
                    values = list(entry[2])
                else:
                    r = entry[1]; values = [getattr(r, c[0], "") for c in cfg["cols"]]
                ws.append([str(v) if v is not None else "" for v in values])
        wb.save(path)

    def gen_sitemap(self):
        if not self.spider or not self.spider.results:
            messagebox.showinfo("SilentRat", "No data yet. Run a crawl first."); return
        path = filedialog.asksaveasfilename(defaultextension=".xml",
            filetypes=[("XML", "*.xml")], initialfile="sitemap.xml")
        if not path:
            return
        n = self.spider.generate_sitemap(path)
        messagebox.showinfo("SilentRat", f"Sitemap generated with {n} URLs:\n{path}")


def main():
    root = tk.Tk()
    try:
        ico = os.path.join(os.path.dirname(os.path.abspath(__file__)), "silentrat.ico")
        if os.path.exists(ico):
            root.iconbitmap(ico)
    except Exception:
        pass
    SilentRatApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
